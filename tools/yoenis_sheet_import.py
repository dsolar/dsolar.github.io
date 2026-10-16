#!/usr/bin/env python3
"""Import past keepers from the league's "Fantasy Keepers" Google Sheet into yoenis/data.json.

Export the sheet as .xlsx (File > Download > Microsoft Excel), then:
  python tools/yoenis_sheet_import.py path/to/Fantasy\\ Keepers.xlsx

Each year tab (2022 on) holds two blocks of eight team columns: a header row of team names, then up to five
keepers written "Player (round/keeper year)", with a * where the commissioner corrected the round. 2022 lists
rounds only for the first block. The sheet's team names aren't always that season's Yahoo names, so each column
is matched to a team by its players, using the keeper picks in that year's draft.

Writes data["keeperHistory"] = {"2026": [[team index, [[player, round, keeper year, corrected], ...]], ...], ...}
and data["offseason"]["sheetOrder"]: the current teams in the sheet's column order, for the 2027 tab.
"""
import json
import re
import sys
import unicodedata
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parent.parent
DATA = ROOT / "yoenis" / "data.json"
CELL = re.compile(r"^(?P<p>.*?)\s*\((?P<r>\d+)(?:/(?P<y>\d))?\)\s*(?P<star>\*)?\s*$")


def fold(s):
    s = unicodedata.normalize("NFKD", s).encode("ascii", "ignore").decode().lower()
    s = re.sub(r"\[(b|p)\]|\b(pitcher|batter)\b|\((p|b)\)", " ", s)
    return re.sub(r"[^a-z0-9 ]", "", re.sub(r"\s+", " ", s)).strip()


def last(s):
    w = fold(s).replace(" jr", "").replace(" ii", "").split()
    return w[-1] if w else ""


def blocks(ws):
    rows = [list(r) for r in ws.iter_rows(values_only=True)]
    out, skip = [], -1
    for i, r in enumerate(rows):
        if i <= skip:
            continue
        cells = [c for c in r[:8]]
        # a header row: eight team names and no "(round)" cells
        if sum(1 for c in cells if isinstance(c, str) and c.strip()) >= 6 and not any(isinstance(c, str) and "(" in c for c in cells) and i + 1 < len(rows):
            nxt = rows[i + 1][:8]
            if not any(isinstance(c, str) and c.strip() for c in nxt):
                continue
            for col in range(8):
                head = cells[col]
                if not isinstance(head, str) or not head.strip():
                    continue
                players = []
                for j in range(i + 1, min(i + 6, len(rows))):
                    v = rows[j][col] if col < len(rows[j]) else None
                    if isinstance(v, str) and v.strip() and v.strip() not in (",",):
                        players.append(v.strip())
                out.append((head.strip(), players))
            skip = i + 5
    return out


def parse(cell):
    m = CELL.match(cell)
    if not m:
        return [cell.rstrip("*").strip(), None, None, cell.endswith("*")]
    return [m["p"].strip(), int(m["r"]), int(m["y"]) if m["y"] else None, bool(m["star"])]


def main(xlsx):
    d = json.loads(DATA.read_text(encoding="utf-8"))
    seasons = {s["year"]: s for s in d["seasons"]}
    wb = openpyxl.load_workbook(xlsx, data_only=True)
    hist, notes = {}, []
    for name in wb.sheetnames:
        if not re.fullmatch(r"20\d\d", name) or int(name) not in seasons:
            continue
        y, s = int(name), seasons[int(name)]
        cols = blocks(wb[name])
        if len(cols) < 12:
            continue
        drafted = {}
        for p in s["draft"]:
            drafted.setdefault(fold(p[2]), []).append((p[1], bool(p[3] if len(p) > 3 else 0)))
        teams = []
        for head, players in cols:
            score = {}
            for raw in players:
                nm = fold(parse(raw)[0])
                hits = drafted.get(nm) or [v for k, vs in drafted.items() if last(k) == last(nm) and k[:1] == nm[:1] for v in vs]
                for ti, kept in hits:
                    score[ti] = score.get(ti, 0) + (2 if kept else 1)
            if head in s["teams"]:
                score[s["teams"].index(head)] = score.get(s["teams"].index(head), 0) + 1.5
            ti = max(score, key=score.get) if score else None
            if ti is None:
                notes.append(f"{y}: couldn't place {head}")
                continue
            ks = [parse(c) for c in players]
            for k in ks:
                if k[1] is None:      # 2022's second block lists names only: take the round from the draft
                    for p in s["draft"]:
                        if p[1] == ti and (fold(p[2]) == fold(k[0]) or (last(p[2]) == last(k[0]) and fold(p[2])[:1] == fold(k[0])[:1])):
                            k[1] = p[0]
                            break
            teams.append([ti, ks, head])
        seen = {}
        for t in teams:
            seen.setdefault(t[0], []).append(t[2])
        for ti, heads in seen.items():
            if len(heads) > 1:
                notes.append(f"{y}: {heads} all matched {s['teams'][ti]}")
        hist[str(y)] = [[t[0], t[1]] for t in teams]
        print(f"{y}: {len(teams)} teams, {sum(len(t[1]) for t in teams)} keepers")
        if y == max(seasons):
            d.setdefault("offseason", {})["sheetOrder"] = [s["teams"][t[0]] for t in teams]
    d["keeperHistory"] = hist
    txt = json.dumps(d, ensure_ascii=False, separators=(",", ":"))
    assert "@" not in txt
    DATA.write_text(txt, encoding="utf-8")
    for n in notes:
        print("note:", n)


if __name__ == "__main__":
    main(sys.argv[1])
